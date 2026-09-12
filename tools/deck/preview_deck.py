"""Approximate pptx -> PNG renderer, for checking layout (overflow, overlap, balance).

Not a fidelity renderer. It draws the shape geometry and wraps text with real font
metrics, which is enough to catch the mistakes that matter in a generated deck.
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu

DPI = 110
SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)

FONTS = {
    (False, False): "/System/Library/Fonts/Supplemental/Arial.ttf",
    (True, False): "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    (False, True): "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
    (True, True): "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf",
}
_cache: dict = {}


def font(size_pt, bold=False, italic=False):
    px = max(6, round(size_pt * DPI / 72))
    key = (px, bold, italic)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FONTS[(bold, italic)], px)
    return _cache[key]


def px(emu):
    return round(Emu(int(emu)).inches * DPI)


def rgb(color):
    try:
        if color and color.type is not None:
            return tuple(color.rgb)
    except Exception:
        pass
    return None


def shape_fill(sh):
    try:
        if sh.fill.type is None or str(sh.fill.type) == "MSO_FILL_TYPE.BACKGROUND (5)":
            return None
        return rgb(sh.fill.fore_color)
    except Exception:
        return None


def shape_line(sh):
    try:
        c = rgb(sh.line.color)
        if c is None:
            return None, 0
        w = sh.line.width.pt if sh.line.width else 1.0
        return c, max(1, round(w * DPI / 72))
    except Exception:
        return None, 0


def draw_geometry(img, drw, sh, x, y, w, h):
    name = ""
    try:
        name = str(sh.auto_shape_type)
    except Exception:
        pass
    fill = shape_fill(sh)
    line, lw = shape_line(sh)
    boxxy = [x, y, x + w, y + h]

    if "OVAL" in name:
        drw.ellipse(boxxy, fill=fill, outline=line, width=lw or 1)
    elif "DONUT" in name:
        hole = img.getpixel((min(img.width - 1, x + w // 2), min(img.height - 1, y + h // 2)))
        drw.ellipse(boxxy, fill=fill)
        inset = round(min(w, h) * 0.19)
        drw.ellipse([x + inset, y + inset, x + w - inset, y + h - inset], fill=hole)
    elif "CHEVRON" in name:
        notch = round(h * 0.32)
        pts = [(x, y), (x + w - notch, y), (x + w, y + h // 2),
               (x + w - notch, y + h), (x, y + h), (x + notch, y + h // 2)]
        drw.polygon(pts, fill=fill, outline=line)
        if line and lw > 1:
            drw.line(pts + [pts[0]], fill=line, width=lw)
    elif "RIGHT_ARROW" in name:
        head = round(w * 0.45)
        shaft = round(h * 0.34)
        pts = [(x, y + shaft), (x + w - head, y + shaft), (x + w - head, y),
               (x + w, y + h // 2), (x + w - head, y + h),
               (x + w - head, y + h - shaft), (x, y + h - shaft)]
        drw.polygon(pts, fill=fill, outline=line)
    elif "ROUNDED_RECTANGLE" in name:
        r = max(2, round(min(w, h) * 0.10))
        drw.rounded_rectangle(boxxy, radius=r, fill=fill, outline=line, width=lw or 1)
    else:
        drw.rectangle(boxxy, fill=fill, outline=line, width=lw or 1)


def wrap(text, fnt, max_w, drw):
    out = []
    for hard in text.split("\n"):
        if not hard:
            out.append("")
            continue
        words, line = hard.split(" "), ""
        for word in words:
            trial = f"{line} {word}".strip()
            if drw.textlength(trial, font=fnt) <= max_w or not line:
                line = trial
            else:
                out.append(line)
                line = word
        out.append(line)
    return out


def draw_text(img, drw, sh, x, y, w, h, overflow):
    if not sh.has_text_frame:
        return
    tf = sh.text_frame
    ml = px(tf.margin_left or 0)
    mr = px(tf.margin_right or 0)
    mt = px(tf.margin_top or 0)
    mb = px(tf.margin_bottom or 0)
    inner_w = max(10, w - ml - mr)

    lines = []
    for para in tf.paragraphs:
        runs = [r for r in para.runs if r.text]
        if not runs:
            lines.append((None, "", 0, 0))
            continue
        r = runs[0]
        text = "".join(rr.text for rr in runs)
        sz = (r.font.size.pt if r.font.size else 11)
        fnt = font(sz, bool(r.font.bold), bool(r.font.italic))
        color = rgb(r.font.color) or (31, 41, 51)
        ls = para.line_spacing if isinstance(para.line_spacing, float) else 1.0
        lh = round(sz * DPI / 72 * 1.22 * (ls or 1.0))
        sa = round((para.space_after.pt if para.space_after else 0) * DPI / 72)
        for ln in wrap(text, fnt, inner_w, drw):
            lines.append((fnt, ln, lh, color))
        if sa:
            lines.append((None, "", sa, 0))
        lines[-1] = (lines[-1][0], lines[-1][1], lines[-1][2] + sa, lines[-1][3]) \
            if lines[-1][0] else lines[-1]

    total = sum(l[2] for l in lines)
    anchor = str(tf.vertical_anchor or "")
    if "MIDDLE" in anchor:
        cy = y + mt + max(0, (h - mt - mb - total) // 2)
    elif "BOTTOM" in anchor:
        cy = y + h - mb - total
    else:
        cy = y + mt

    if total > (h - mt - mb) + 3 and any(l[1] for l in lines):
        overflow.append((sh.name, sh.shape_id, total, h - mt - mb))

    align = PP_ALIGN.LEFT
    for para in tf.paragraphs:
        if para.alignment is not None:
            align = para.alignment
            break

    for fnt, ln, lh, color in lines:
        if fnt is None:
            cy += lh
            continue
        tw = drw.textlength(ln, font=fnt)
        if align == PP_ALIGN.CENTER:
            tx = x + ml + (inner_w - tw) / 2
        elif align == PP_ALIGN.RIGHT:
            tx = x + ml + inner_w - tw
        else:
            tx = x + ml
        drw.text((tx, cy), ln, font=fnt, fill=tuple(color))
        cy += lh


def render(slide, index):
    W = round(13.333 * DPI)
    H = round(7.5 * DPI)
    img = Image.new("RGB", (W, H), (255, 255, 255))
    drw = ImageDraw.Draw(img)
    overflow = []
    for sh in slide.shapes:
        x, y = px(sh.left), px(sh.top)
        w, h = max(1, px(sh.width)), max(1, px(sh.height))
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                im = Image.open(sh.image.blob and __import__("io").BytesIO(sh.image.blob))
                img.paste(im.convert("RGB").resize((w, h), Image.LANCZOS), (x, y))
            except Exception as exc:
                drw.rectangle([x, y, x + w, y + h], outline=(200, 0, 0))
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            draw_text(img, drw, sh, x, y, w, h, overflow)
            continue
        rot = getattr(sh, "rotation", 0) or 0
        if abs(rot) > 0.5:
            layer = Image.new("RGBA", (w * 3, h * 3), (0, 0, 0, 0))
            ld = ImageDraw.Draw(layer)
            fill = shape_fill(sh)
            ld.rectangle([w, h, w * 2, h * 2], fill=tuple(fill) if fill else None)
            layer = layer.rotate(-rot, resample=Image.BICUBIC, center=(w * 1.5, h * 1.5))
            img.paste(layer, (x - w, y - h), layer)
            continue
        draw_geometry(img, drw, sh, x, y, w, h)
        draw_text(img, drw, sh, x, y, w, h, overflow)
    img.save(OUT / f"slide{index}.png")
    return overflow


prs = Presentation(SRC)
for i, slide in enumerate(prs.slides, 1):
    over = render(slide, i)
    if over:
        print(f"slide {i}: {len(over)} text overflow(s)")
        for name, sid, need, have in over:
            print(f"    {name} (id {sid}): needs {need}px, has {have}px")
    else:
        print(f"slide {i}: ok")
